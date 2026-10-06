"""レポートをExcel(.xlsx)に書き出す。"""

import io

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

_HEADER_FILL = PatternFill("solid", fgColor="DDE7F3")
_MAX_WIDTH = 60


def to_excel_bytes(sheets: dict[str, pd.DataFrame]) -> bytes:
    buffer = io.BytesIO()
    with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
        for name, frame in sheets.items():
            frame.to_excel(writer, sheet_name=name[:31], index=False)
            sheet = writer.sheets[name[:31]]
            sheet.freeze_panes = "A2"
            for cell in sheet[1]:
                cell.font = Font(bold=True)
                cell.fill = _HEADER_FILL
            for i, column in enumerate(frame.columns, 1):
                values = [str(column), *frame[column].astype(str).tolist()]
                width = min(max(len(v) for v in values) * 1.8 + 2, _MAX_WIDTH)
                sheet.column_dimensions[get_column_letter(i)].width = width
                if width >= _MAX_WIDTH:
                    for cell in sheet[get_column_letter(i)][1:]:
                        cell.alignment = Alignment(wrap_text=True, vertical="top")
    return buffer.getvalue()
